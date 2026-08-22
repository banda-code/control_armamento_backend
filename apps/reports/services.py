from collections import defaultdict
from io import BytesIO

from django.core.exceptions import ValidationError
from django.db.models import Count, Sum
from django.utils import timezone

from apps.accounts.scopes import get_user_unit_id
from apps.assignments.models import (
    AssignmentStatus,
    IndividualAssignment,
)
from apps.personnel.models import Personnel
from apps.inventory.models import (
    ControlMethod,
    MaterialSpecification,
    SerializedMaterial,
    StockBatch,
)
from apps.organization.models import Unit
from apps.reports.models import IndividualFiliationConfig


CATEGORY_ORDER = {
    "ARMAMENTO": 10,
    "ARMA BLANCA": 20,
    "CLASE V": 30,
    "MUNICION": 30,
    "MUNICIÓN": 30,
    "AGENTES QUIMICOS": 40,
    "AGENTES QUÍMICOS": 40,
    "EXPLOSIVOS": 50,
}


# ============================================================
# FILIACIÓN / REGISTRO DE ARMAMENTO DE DOTACIÓN INDIVIDUAL
# ============================================================


def build_individual_filiation_report(
    unit,
    personnel_id,
):
    """
    Construye los datos necesarios para generar
    el Registro de Armamento de Dotación Individual.
    """

    # --------------------------------------------------------
    # PERSONAL
    # --------------------------------------------------------

    try:
        personnel = (
            Personnel.objects
            .select_related(
                "rank",
                "position",
                "unit",
                "section",
            )
            .get(
                pk=personnel_id,
                unit=unit,
                is_active=True,
            )
        )
    except Personnel.DoesNotExist as exc:
        raise ValidationError(
            "El personal seleccionado no existe "
            "o no pertenece a la unidad."
        ) from exc

    # --------------------------------------------------------
    # DOTACIONES INDIVIDUALES ACTIVAS
    # --------------------------------------------------------

    assignments = list(
        IndividualAssignment.objects
        .filter(
            personnel=personnel,
            status=AssignmentStatus.ACTIVE,
            material__is_active=True,
        )
        .select_related(
            "material",
            "material__unit",
            "material__armory",
            "material__specification",
            "material__specification__material_type",
            "material__specification__manufacturer",
            "material__specification__country",
            "material__specification__caliber",
        )
        .prefetch_related(
            "assigned_components__component__component_type",
            "photos",
        )
        .order_by(
            "material__specification__material_type__name",
            "assigned_at",
        )
    )

    if not assignments:
        raise ValidationError(
            "El personal seleccionado no tiene "
            "dotaciones individuales activas."
        )

    # --------------------------------------------------------
    # MATERIALES
    # --------------------------------------------------------

    materials = []

    for assignment in assignments:
        material = assignment.material
        specification = material.specification
        material_type = specification.material_type

        components = []

        magazine_count = 0

        for assigned_component in (
            assignment.assigned_components.all()
        ):
            component = assigned_component.component
            component_type = component.component_type

            quantity = (
                assigned_component.quantity_delivered
            )

            components.append(
                {
                    "type": component_type.name,
                    "identification_number": (
                        component.identification_number
                        or ""
                    ),
                    "quantity": quantity,
                    "observations": (
                        assigned_component.observations
                        or ""
                    ),
                }
            )

            if (
                "CARGADOR"
                in component_type.name.upper()
            ):
                magazine_count += int(quantity or 0)

        material_data = {
            "assignment_id": str(assignment.id),

            "material_type": material_type.name,

            "institutional_code": (
                material.institutional_code
            ),

            "identification_number": (
                material.identification_number
            ),

            "model": specification.name,

            "manufacturer": (
                specification.manufacturer.name
                if specification.manufacturer
                else ""
            ),

            "country": (
                specification.country.name
                if specification.country
                else ""
            ),

            "caliber": (
                specification.caliber.name
                if specification.caliber
                else ""
            ),

            "observations": (
                material.observations
                or assignment.assignment_observations
                or ""
            ),

            "magazine_count": magazine_count,

            "components": components,

            "photos": list(
                assignment.photos.all()
            ),
        }

        materials.append(material_data)

    # --------------------------------------------------------
    # IDENTIFICAR PISTOLA Y CUCHILLO BAYONETA
    # --------------------------------------------------------

    pistol = None
    bayonet = None

    for material in materials:
        material_type_name = (
            material["material_type"]
            .strip()
            .upper()
        )

        if (
            pistol is None
            and "PISTOLA" in material_type_name
        ):
            pistol = material

        if (
            bayonet is None
            and (
                "BAYONETA" in material_type_name
                or "CUCHILLO" in material_type_name
            )
        ):
            bayonet = material

    # --------------------------------------------------------
    # RESULTADO
    # --------------------------------------------------------

    # ========================================================
    # FIRMANTES CONFIGURADOS PARA LA UNIDAD
    # ========================================================

    filiation_config = (
        IndividualFiliationConfig.objects
        .filter(unit=unit)
        .select_related(
            "verification_responsible__rank",
            "verification_responsible__position",
            "logistics_chief__rank",
            "logistics_chief__position",
            "unit_commander__rank",
            "unit_commander__position",
        )
        .first()
    )

    def signer_data(person):
        if not person:
            return None

        return {
            "id": str(person.id),
            "full_name": person.full_name,
            "rank": (
                person.rank.abbreviation
                if person.rank
                else ""
            ),
            "position": (
                person.position.name
                if person.position
                else ""
            ),
        }

    signatories = {
        "verification_responsible": signer_data(
            filiation_config.verification_responsible
            if filiation_config
            else None
        ),
        "logistics_chief": signer_data(
            filiation_config.logistics_chief
            if filiation_config
            else None
        ),
        "unit_commander": signer_data(
            filiation_config.unit_commander
            if filiation_config
            else None
        ),
    }

    return {
        "unit": {
            "id": str(unit.id),
            "name": unit.name,
        },

        "personnel": {
            "id": str(personnel.id),
            "tin": personnel.tin,

            "full_name": personnel.full_name,

            "rank": (
                personnel.rank.name
                if personnel.rank
                else ""
            ),

            "rank_abbreviation": (
                personnel.rank.abbreviation
                if personnel.rank
                else ""
            ),

            "graduation_year": (
                personnel.graduation_year
            ),

            "position": (
                personnel.position.name
                if personnel.position
                else ""
            ),

            "section": (
                personnel.section.name
                if personnel.section
                else ""
            ),
        },

        "pistol": pistol,
        "bayonet": bayonet,
        "materials": materials,
        "total_assignments": len(materials),
        "signatories": signatories,
    }


# ============================================================
# ÁMBITO DE UNIDAD
# ============================================================


def resolve_report_unit(user, requested_unit_id=None):
    """
    Devuelve la unidad sobre la que puede emitirse el reporte.

    - Usuario con alcance global: puede indicar ?unit=<uuid>.
    - Usuario sin alcance global: siempre queda limitado a su unidad.
    """

    if getattr(user, "has_global_scope", False):
        unit_id = requested_unit_id or getattr(user, "unit_id", None)

        if not unit_id:
            raise ValidationError(
                "Debe seleccionar una unidad para generar el reporte."
            )

        try:
            return Unit.objects.get(pk=unit_id)
        except Unit.DoesNotExist as exc:
            raise ValidationError("La unidad seleccionada no existe.") from exc

    unit_id = get_user_unit_id(user)

    if not unit_id:
        raise ValidationError(
            "El usuario no tiene una unidad asignada para generar reportes."
        )

    try:
        return Unit.objects.get(pk=unit_id)
    except Unit.DoesNotExist as exc:
        raise ValidationError("La unidad del usuario no existe.") from exc


# ============================================================
# OPCIONES DE ARMAMENTO / MATERIAL SERIALIZADO
# ============================================================


def get_armament_options(unit):
    rows = (
        SerializedMaterial.objects
        .filter(
            unit=unit,
            is_active=True,
            specification__is_active=True,
            specification__material_type__is_active=True,
            specification__material_type__control_method=(
                ControlMethod.SERIALIZED
            ),
        )
        .values(
            "specification_id",
            "specification__material_type__category__name",
            "specification__material_type__name",
            "specification__name",
            "specification__manufacturer__name",
            "specification__country__name",
            "specification__caliber__name",
        )
        .annotate(total=Count("id"))
        .order_by(
            "specification__material_type__category__name",
            "specification__material_type__name",
            "specification__name",
        )
    )

    return [
        {
            "id": str(row["specification_id"]),
            "category": (
                row["specification__material_type__category__name"] or ""
            ),
            "material_type": row["specification__material_type__name"] or "",
            "specification": row["specification__name"] or "",
            "manufacturer": row["specification__manufacturer__name"] or "",
            "country": row["specification__country__name"] or "",
            "caliber": row["specification__caliber__name"] or "",
            "total": row["total"],
            "label": " - ".join(
                value
                for value in [
                    row["specification__material_type__name"],
                    row["specification__name"],
                    row["specification__caliber__name"],
                ]
                if value
            ),
        }
        for row in rows
    ]


# ============================================================
# REPORTE POR TIPO / ESPECIFICACIÓN DE ARMAMENTO
# ============================================================


def build_armament_detail(unit, specification_id):
    try:
        specification = (
            MaterialSpecification.objects
            .select_related(
                "material_type",
                "material_type__category",
                "manufacturer",
                "country",
                "caliber",
            )
            .get(pk=specification_id)
        )
    except MaterialSpecification.DoesNotExist as exc:
        raise ValidationError("El armamento seleccionado no existe.") from exc

    materials = list(
        SerializedMaterial.objects
        .filter(
            unit=unit,
            specification=specification,
            is_active=True,
        )
        .select_related(
            "unit",
            "armory",
            "specification",
            "specification__material_type",
            "specification__manufacturer",
            "specification__country",
            "specification__caliber",
        )
        .prefetch_related(
            "components__component_type",
        )
        .order_by(
            "institutional_code",
            "identification_number",
        )
    )

    status_summary = defaultdict(int)
    condition_summary = defaultdict(int)
    rows = []

    for index, material in enumerate(materials, start=1):
        status_summary[material.get_status_display()] += 1
        condition_summary[material.get_physical_condition_display()] += 1

        components = []
        for component in material.components.all():
            if not component.is_active:
                continue

            description = component.component_type.name

            if component.identification_number:
                description += f" ({component.identification_number})"

            if component.quantity > 1:
                description += f" x{component.quantity}"

            components.append(description)

        rows.append(
            {
                "nro": index,
                "institutional_code": material.institutional_code,
                "identification_number": material.identification_number,
                "status": material.status,
                "status_display": material.get_status_display(),
                "physical_condition": material.physical_condition,
                "physical_condition_display": (
                    material.get_physical_condition_display()
                ),
                "armory": material.armory.name,
                "manufacturing_year": material.manufacturing_year or "",
                "acquisition_source": material.acquisition_source or "",
                "registration_document": (
                    material.registration_document or ""
                ),
                "components": ", ".join(components),
                "observations": material.observations or "",
            }
        )

    return {
        "report_type": "ARMAMENT_DETAIL",
        "generated_at": timezone.localtime(),
        "unit": {
            "id": str(unit.id),
            "name": unit.name,
        },
        "weapon": {
            "id": str(specification.id),
            "category": specification.material_type.category.name,
            "material_type": specification.material_type.name,
            "specification": specification.name,
            "manufacturer": (
                specification.manufacturer.name
                if specification.manufacturer
                else ""
            ),
            "country": (
                specification.country.name
                if specification.country
                else ""
            ),
            "caliber": (
                specification.caliber.name
                if specification.caliber
                else ""
            ),
            "description": specification.description or "",
        },
        "total": len(rows),
        "status_summary": dict(status_summary),
        "condition_summary": dict(condition_summary),
        "rows": rows,
    }


# ============================================================
# REPORTE GENERAL DE MATERIAL BÉLICO
# ============================================================


def build_general_material_report(unit):
    grouped = defaultdict(list)

    serialized_rows = (
        SerializedMaterial.objects
        .filter(
            unit=unit,
            is_active=True,
            specification__is_active=True,
        )
        .values(
            "specification_id",
            "specification__material_type__category__name",
            "specification__material_type__name",
            "specification__name",
            "specification__caliber__name",
        )
        .annotate(total=Count("id"))
        .order_by(
            "specification__material_type__category__name",
            "specification__material_type__name",
            "specification__name",
        )
    )

    for row in serialized_rows:
        category = (
            row["specification__material_type__category__name"]
            or "SIN CATEGORÍA"
        )

        description = " ".join(
            value
            for value in [
                row["specification__material_type__name"],
                row["specification__name"],
                row["specification__caliber__name"],
            ]
            if value
        )

        grouped[category].append(
            {
                "source": "SERIALIZED",
                "specification_id": str(row["specification_id"]),
                "description": description,
                "unit_of_measure": "UND.",
                "total": row["total"],
            }
        )

    quantity_rows = (
        StockBatch.objects
        .filter(
            unit=unit,
            is_active=True,
            stock_material__is_active=True,
            stock_material__specification__is_active=True,
        )
        .values(
            "stock_material__specification_id",
            "stock_material__specification__material_type__category__name",
            "stock_material__specification__material_type__name",
            "stock_material__specification__name",
            "stock_material__specification__caliber__name",
            "stock_material__unit_of_measure__symbol",
            "stock_material__unit_of_measure__name",
        )
        .annotate(total=Sum("current_quantity"))
        .order_by(
            "stock_material__specification__material_type__category__name",
            "stock_material__specification__material_type__name",
            "stock_material__specification__name",
        )
    )

    for row in quantity_rows:
        category = (
            row[
                "stock_material__specification__material_type__category__name"
            ]
            or "SIN CATEGORÍA"
        )

        description = " ".join(
            value
            for value in [
                row["stock_material__specification__material_type__name"],
                row["stock_material__specification__name"],
                row["stock_material__specification__caliber__name"],
            ]
            if value
        )

        grouped[category].append(
            {
                "source": "QUANTITY",
                "specification_id": str(
                    row["stock_material__specification_id"]
                ),
                "description": description,
                "unit_of_measure": (
                    row["stock_material__unit_of_measure__symbol"]
                    or row["stock_material__unit_of_measure__name"]
                    or "UND."
                ),
                "total": float(row["total"] or 0),
            }
        )

    sections = []

    for category, rows in grouped.items():
        subtotal_by_unit = defaultdict(float)

        for row in rows:
            subtotal_by_unit[row["unit_of_measure"]] += float(row["total"])

        sections.append(
            {
                "category": category,
                "rows": sorted(
                    rows,
                    key=lambda item: item["description"],
                ),
                "subtotal_by_unit": dict(subtotal_by_unit),
            }
        )

    sections.sort(
        key=lambda section: (
            CATEGORY_ORDER.get(section["category"].upper(), 999),
            section["category"],
        )
    )

    return {
        "report_type": "GENERAL_MATERIAL",
        "generated_at": timezone.localtime(),
        "unit": {
            "id": str(unit.id),
            "name": unit.name,
        },
        "sections": sections,
        "total_sections": len(sections),
    }


# ============================================================
# EXCEL
# ============================================================


def _excel_styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    navy = "0F172A"
    slate = "334155"
    light = "E2E8F0"
    white = "FFFFFF"

    thin = Side(style="thin", color="94A3B8")

    return {
        "navy": navy,
        "slate": slate,
        "light": light,
        "white": white,
        "title_font": Font(name="Arial", bold=True, size=14, color=navy),
        "subtitle_font": Font(name="Arial", bold=True, size=11, color=navy),
        "normal_font": Font(name="Arial", size=9, color="111827"),
        "header_font": Font(name="Arial", bold=True, size=9, color=white),
        "section_font": Font(name="Arial", bold=True, size=10, color=white),
        "navy_fill": PatternFill("solid", fgColor=navy),
        "slate_fill": PatternFill("solid", fgColor=slate),
        "light_fill": PatternFill("solid", fgColor=light),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "left": Alignment(horizontal="left", vertical="center", wrap_text=True),
    }


def export_armament_detail_xlsx(report, generated_by=""):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    styles = _excel_styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "Armamento"

    ws.merge_cells("A1:K1")
    ws["A1"] = "ARMADA BOLIVIANA"
    ws["A1"].font = styles["title_font"]
    ws["A1"].alignment = styles["center"]

    ws.merge_cells("A2:K2")
    ws["A2"] = report["unit"]["name"]
    ws["A2"].font = styles["subtitle_font"]
    ws["A2"].alignment = styles["center"]

    ws.merge_cells("A4:K4")
    ws["A4"] = "INVENTARIO DE MATERIAL BÉLICO - DETALLE POR ARMAMENTO"
    ws["A4"].font = styles["subtitle_font"]
    ws["A4"].alignment = styles["center"]

    weapon = report["weapon"]

    metadata = [
        ("TIPO DE ARMAMENTO", weapon["material_type"]),
        ("MODELO / ESPECIFICACIÓN", weapon["specification"]),
        ("FABRICANTE", weapon["manufacturer"] or "S/D"),
        ("INDUSTRIA / PAÍS", weapon["country"] or "S/D"),
        ("CALIBRE", weapon["caliber"] or "S/D"),
        ("TOTAL EN LA UNIDAD", report["total"]),
    ]

    row = 6
    for label, value in metadata:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=1).font = styles["header_font"]
        ws.cell(row=row, column=1).fill = styles["slate_fill"]
        ws.cell(row=row, column=1).alignment = styles["left"]

        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=11)
        ws.cell(row=row, column=4, value=value)
        ws.cell(row=row, column=4).font = styles["normal_font"]
        ws.cell(row=row, column=4).alignment = styles["left"]
        row += 1

    row += 1

    headers = [
        "N°",
        "CÓDIGO",
        "N° SERIE / IDENTIFICACIÓN",
        "ESTADO",
        "CONDICIÓN",
        "PAÑOL",
        "AÑO FAB.",
        "PROCEDENCIA / ADQUISICIÓN",
        "DOCUMENTO DE ALTA",
        "ACCESORIOS / COMPONENTES",
        "OBSERVACIONES",
    ]

    header_row = row

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = styles["header_font"]
        cell.fill = styles["navy_fill"]
        cell.alignment = styles["center"]
        cell.border = styles["border"]

    for item in report["rows"]:
        row += 1
        values = [
            item["nro"],
            item["institutional_code"],
            item["identification_number"],
            item["status_display"],
            item["physical_condition_display"],
            item["armory"],
            item["manufacturing_year"],
            item["acquisition_source"],
            item["registration_document"],
            item["components"],
            item["observations"],
        ]

        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = styles["normal_font"]
            cell.alignment = styles["center"] if col in {1, 4, 5, 7} else styles["left"]
            cell.border = styles["border"]

    row += 2
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.cell(row=row, column=1, value=f"TOTAL: {report['total']} REGISTROS")
    ws.cell(row=row, column=1).font = styles["subtitle_font"]

    row += 2
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
    ws.cell(
        row=row,
        column=1,
        value=(
            f"Generado por: {generated_by or '—'} | "
            f"Fecha: {report['generated_at'].strftime('%d/%m/%Y %H:%M')}"
        ),
    )
    ws.cell(row=row, column=1).font = styles["normal_font"]
    ws.cell(row=row, column=1).alignment = styles["left"]

    widths = [6, 16, 22, 16, 16, 22, 12, 26, 24, 30, 34]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:K{max(header_row, row - 4)}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"1:{header_row}"

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def export_general_material_xlsx(report, generated_by=""):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    styles = _excel_styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "Parte General"

    ws.merge_cells("A1:D1")
    ws["A1"] = "ARMADA BOLIVIANA"
    ws["A1"].font = styles["title_font"]
    ws["A1"].alignment = styles["center"]

    ws.merge_cells("A2:D2")
    ws["A2"] = report["unit"]["name"]
    ws["A2"].font = styles["subtitle_font"]
    ws["A2"].alignment = styles["center"]

    ws.merge_cells("A4:D4")
    ws["A4"] = "PARTE GENERAL DE MATERIAL BÉLICO Y EXISTENCIAS"
    ws["A4"].font = styles["subtitle_font"]
    ws["A4"].alignment = styles["center"]

    row = 6

    for index, section in enumerate(report["sections"], start=1):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.cell(
            row=row,
            column=1,
            value=f"{index}.- {section['category']}",
        )
        ws.cell(row=row, column=1).font = styles["section_font"]
        ws.cell(row=row, column=1).fill = styles["slate_fill"]
        ws.cell(row=row, column=1).alignment = styles["left"]
        row += 1

        headers = ["N°", "TIPO / ESPECIFICACIÓN", "UNIDAD MEDIDA", "TOTAL"]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = styles["header_font"]
            cell.fill = styles["navy_fill"]
            cell.alignment = styles["center"]
            cell.border = styles["border"]

        for item_index, item in enumerate(section["rows"], start=1):
            row += 1
            values = [
                item_index,
                item["description"],
                item["unit_of_measure"],
                item["total"],
            ]

            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.font = styles["normal_font"]
                cell.alignment = styles["center"] if col in {1, 3, 4} else styles["left"]
                cell.border = styles["border"]

        for unit_symbol, subtotal in section["subtotal_by_unit"].items():
            row += 1
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            ws.cell(row=row, column=1, value=f"SUBTOTAL {unit_symbol}")
            ws.cell(row=row, column=1).font = styles["subtitle_font"]
            ws.cell(row=row, column=1).fill = styles["light_fill"]
            ws.cell(row=row, column=3, value=unit_symbol)
            ws.cell(row=row, column=4, value=subtotal)

            for col in range(1, 5):
                ws.cell(row=row, column=col).border = styles["border"]
                ws.cell(row=row, column=col).alignment = styles["center"]

        row += 2

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.cell(
        row=row,
        column=1,
        value=(
            f"Generado por: {generated_by or '—'} | "
            f"Fecha: {report['generated_at'].strftime('%d/%m/%Y %H:%M')}"
        ),
    )
    ws.cell(row=row, column=1).font = styles["normal_font"]

    widths = [7, 80, 18, 18]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ============================================================
# PDF
# ============================================================


def _pdf_base_styles():
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    base = getSampleStyleSheet()

    return {
        "title": ParagraphStyle(
            "ReportTitle",
            parent=base["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=14,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#0F172A"),
            spaceAfter=4,
        ),
        "subtitle": ParagraphStyle(
            "ReportSubtitle",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=9,
            leading=11,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#334155"),
            spaceAfter=3,
        ),
        "normal": ParagraphStyle(
            "ReportNormal",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=7.5,
            leading=9,
            alignment=TA_LEFT,
        ),
        "small": ParagraphStyle(
            "ReportSmall",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=6.5,
            leading=8,
            alignment=TA_LEFT,
        ),
    }


def export_armament_detail_pdf(report, generated_by=""):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = _pdf_base_styles()
    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=8 * mm,
        leftMargin=8 * mm,
        topMargin=8 * mm,
        bottomMargin=8 * mm,
        title="Reporte de Armamento - FORTALEZA",
        author="Sistema FORTALEZA",
    )

    story = []
    story.append(Paragraph("ARMADA BOLIVIANA", styles["title"]))
    story.append(Paragraph(report["unit"]["name"], styles["subtitle"]))
    story.append(
        Paragraph(
            "INVENTARIO DE MATERIAL BÉLICO - DETALLE POR ARMAMENTO",
            styles["subtitle"],
        )
    )
    story.append(Spacer(1, 3 * mm))

    weapon = report["weapon"]

    meta_data = [
        ["TIPO", weapon["material_type"], "MODELO", weapon["specification"]],
        ["FABRICANTE", weapon["manufacturer"] or "S/D", "PAÍS", weapon["country"] or "S/D"],
        ["CALIBRE", weapon["caliber"] or "S/D", "TOTAL", str(report["total"])],
    ]

    meta_table = Table(
        meta_data,
        colWidths=[28 * mm, 66 * mm, 28 * mm, 66 * mm],
    )
    meta_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E2E8F0")),
                ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#E2E8F0")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(meta_table)
    story.append(Spacer(1, 4 * mm))

    headers = [
        "N°",
        "CÓDIGO",
        "SERIE / IDENT.",
        "ESTADO",
        "CONDICIÓN",
        "PAÑOL",
        "AÑO",
        "PROCEDENCIA",
        "DOCUMENTO",
        "ACCESORIOS",
        "OBSERVACIONES",
    ]

    table_data = [headers]

    for item in report["rows"]:
        table_data.append(
            [
                str(item["nro"]),
                item["institutional_code"],
                item["identification_number"],
                item["status_display"],
                item["physical_condition_display"],
                item["armory"],
                str(item["manufacturing_year"] or ""),
                Paragraph(item["acquisition_source"] or "", styles["small"]),
                Paragraph(item["registration_document"] or "", styles["small"]),
                Paragraph(item["components"] or "", styles["small"]),
                Paragraph(item["observations"] or "", styles["small"]),
            ]
        )

    detail_table = Table(
        table_data,
        repeatRows=1,
        colWidths=[
            7 * mm,
            19 * mm,
            25 * mm,
            19 * mm,
            19 * mm,
            24 * mm,
            11 * mm,
            29 * mm,
            28 * mm,
            35 * mm,
            43 * mm,
        ],
    )

    detail_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 6.5),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 1), (-1, -1), 6.2),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#94A3B8")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 2.5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2.5),
                ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ]
        )
    )

    story.append(detail_table)
    story.append(Spacer(1, 4 * mm))
    story.append(
        Paragraph(
            (
                f"Total registrado: {report['total']} | "
                f"Generado por: {generated_by or '—'} | "
                f"Fecha: {report['generated_at'].strftime('%d/%m/%Y %H:%M')}"
            ),
            styles["normal"],
        )
    )

    doc.build(story)
    buffer.seek(0)
    return buffer


def export_general_material_pdf(report, generated_by=""):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = _pdf_base_styles()
    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title="Parte General de Material Bélico - FORTALEZA",
        author="Sistema FORTALEZA",
    )

    story = []
    story.append(Paragraph("ARMADA BOLIVIANA", styles["title"]))
    story.append(Paragraph(report["unit"]["name"], styles["subtitle"]))
    story.append(
        Paragraph(
            "PARTE GENERAL DE MATERIAL BÉLICO Y EXISTENCIAS",
            styles["subtitle"],
        )
    )
    story.append(Spacer(1, 4 * mm))

    for index, section in enumerate(report["sections"], start=1):
        section_header = Table(
            [[f"{index}.- {section['category']}"]],
            colWidths=[180 * mm],
        )
        section_header.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#334155")),
                    ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
                    ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(section_header)

        table_data = [["N°", "TIPO / ESPECIFICACIÓN", "U.M.", "TOTAL"]]

        for item_index, item in enumerate(section["rows"], start=1):
            table_data.append(
                [
                    str(item_index),
                    Paragraph(item["description"], styles["normal"]),
                    item["unit_of_measure"],
                    str(item["total"]),
                ]
            )

        for symbol, subtotal in section["subtotal_by_unit"].items():
            table_data.append(
                [
                    "",
                    Paragraph(f"<b>SUBTOTAL {symbol}</b>", styles["normal"]),
                    symbol,
                    str(subtotal),
                ]
            )

        table = Table(
            table_data,
            repeatRows=1,
            colWidths=[12 * mm, 128 * mm, 20 * mm, 20 * mm],
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 7),
                    ("ALIGN", (0, 0), (0, -1), "CENTER"),
                    ("ALIGN", (2, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#94A3B8")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        story.append(table)
        story.append(Spacer(1, 5 * mm))

    story.append(
        Paragraph(
            (
                f"Generado por: {generated_by or '—'} | "
                f"Fecha: {report['generated_at'].strftime('%d/%m/%Y %H:%M')}"
            ),
            styles["normal"],
        )
    )

    doc.build(story)
    buffer.seek(0)
    return buffer


# ============================================================
# PDF: FILIACIÓN / REGISTRO DE ARMAMENTO DE DOTACIÓN INDIVIDUAL
# ============================================================


# ============================================================
# PDF: REGISTRO DE ARMAMENTO DE DOTACIÓN INDIVIDUAL
# ============================================================


def export_individual_filiation_pdf(
    report,
    generated_by="",
):
    """
    Genera la filiación de armamento individual
    en formato institucional HOJA CARTA.
    """

    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    stream = BytesIO()

    page_width, page_height = LETTER

    pdf = canvas.Canvas(
        stream,
        pagesize=LETTER,
    )

    pdf.setTitle(
        "Registro de Armamento de Dotación Individual"
    )

    # ========================================================
    # UTILIDADES
    # ========================================================

    def safe(value, default=""):
        if value in (None, ""):
            return default

        return str(value)

    def draw_field(
        label,
        value,
        x,
        y,
        label_width=48 * mm,
        font_size=9.5,
    ):
        pdf.setFont(
            "Helvetica-Bold",
            font_size,
        )

        pdf.drawString(
            x,
            y,
            label,
        )

        pdf.drawString(
            x + label_width - 4 * mm,
            y,
            ":",
        )

        pdf.setFont(
            "Helvetica",
            font_size,
        )

        pdf.drawString(
            x + label_width,
            y,
            safe(value),
        )

    def get_photo_reader(photo):
        image_field = getattr(
            photo,
            "photo",
            None,
        )

        if not image_field:
            return None

        try:
            return ImageReader(
                image_field.path
            )
        except Exception:
            pass

        try:
            image_field.open("rb")

            return ImageReader(
                image_field.file
            )
        except Exception:
            return None

    def draw_photo_box(
        photo,
        x,
        y,
        width,
        height,
    ):
        pdf.setLineWidth(1)

        pdf.rect(
            x,
            y,
            width,
            height,
        )

        reader = (
            get_photo_reader(photo)
            if photo
            else None
        )

        if not reader:
            pdf.setFont(
                "Helvetica",
                7,
            )

            pdf.drawCentredString(
                x + width / 2,
                y + height / 2,
                "SIN FOTOGRAFÍA",
            )

            return

        try:
            image_width, image_height = (
                reader.getSize()
            )

            available_width = (
                width - 1.5 * mm
            )

            available_height = (
                height - 1.5 * mm
            )

            scale = min(
                available_width / image_width,
                available_height / image_height,
            )

            draw_width = (
                image_width * scale
            )

            draw_height = (
                image_height * scale
            )

            draw_x = (
                x
                + (
                    width
                    - draw_width
                ) / 2
            )

            draw_y = (
                y
                + (
                    height
                    - draw_height
                ) / 2
            )

            pdf.drawImage(
                reader,
                draw_x,
                draw_y,
                width=draw_width,
                height=draw_height,
                preserveAspectRatio=True,
                mask="auto",
            )

        except Exception:
            pdf.setFont(
                "Helvetica",
                7,
            )

            pdf.drawCentredString(
                x + width / 2,
                y + height / 2,
                "FOTOGRAFÍA NO DISPONIBLE",
            )

    # ========================================================
    # DATOS
    # ========================================================

    unit = report.get(
        "unit",
        {},
    )

    personnel = report.get(
        "personnel",
        {},
    )

    pistol = (
        report.get("pistol")
        or {}
    )

    bayonet = (
        report.get("bayonet")
        or {}
    )

    # ========================================================
    # FOTOGRAFÍAS
    #
    # 1 GENERAL
    # 2 COMPONENTS
    # 3 SERIAL
    # ========================================================

    photo_priority = {
        "GENERAL": 10,
        "COMPONENTS": 20,
        "SERIAL": 30,
        "LEFT_SIDE": 40,
        "RIGHT_SIDE": 50,
    }

    photos_found = []
    seen_photo_ids = set()

    for material in report.get(
        "materials",
        [],
    ):
        for photo in material.get(
            "photos",
            [],
        ):
            if getattr(
                photo,
                "moment",
                "",
            ) != "DELIVERY":
                continue

            photo_id = str(
                getattr(
                    photo,
                    "id",
                    id(photo),
                )
            )

            if photo_id in seen_photo_ids:
                continue

            seen_photo_ids.add(
                photo_id
            )

            photos_found.append(
                photo
            )

    photos_found.sort(
        key=lambda photo: (
            photo_priority.get(
                getattr(
                    photo,
                    "photo_type",
                    "",
                ),
                99,
            ),
            str(
                getattr(
                    photo,
                    "created_at",
                    "",
                )
            ),
        )
    )

    photos = photos_found[:3]

    while len(photos) < 3:
        photos.append(None)

    # ========================================================
    # MEDIDAS GENERALES
    # ========================================================

    left_margin = 28 * mm
    right_margin = 18 * mm

    photo_width = 50 * mm
    photo_x = (
        page_width
        - right_margin
        - photo_width
    )

    text_right = (
        photo_x - 7 * mm
    )

    # Centro del bloque izquierdo.
    # Esto centra ARMADA BOLIVIANA,
    # BATALLÓN... y BOLIVIA tal como el Word.
    header_center_x = 55 * mm

    # ========================================================
    # MEMBRETE
    # ========================================================

    pdf.setFont(
        "Helvetica-Bold",
        9,
    )

    pdf.drawCentredString(
        header_center_x,
        page_height - 17 * mm,
        "ARMADA BOLIVIANA",
    )

    pdf.drawCentredString(
        header_center_x,
        page_height - 22 * mm,
        safe(
            unit.get("name"),
            "BATALLÓN DE COMANDOS ANFIBIOS",
        ).upper(),
    )

    pdf.drawCentredString(
        header_center_x,
        page_height - 27 * mm,
        "BOLIVIA",
    )

    # ========================================================
    # TÍTULO
    # ========================================================

    pdf.setFont(
        "Helvetica-Bold",
        12,
    )

    pdf.drawCentredString(
        page_width / 2,
        page_height - 38 * mm,
        (
            "REGISTRO DE ARMAMENTO "
            "DE DOTACIÓN INDIVIDUAL"
        ),
    )

    # ========================================================
    # MARCA DE AGUA BCA
    # ========================================================

    from pathlib import Path

    watermark_path = (
        Path(__file__).resolve().parent
        / "assets"
        / "bca_watermark.png"
    )

    if watermark_path.exists():
        try:
            watermark_width = 92 * mm
            watermark_height = 92 * mm

            watermark_x = 65 * mm
            watermark_y = 95 * mm

            pdf.saveState()

            pdf.drawImage(
                str(watermark_path),
                watermark_x,
                watermark_y,
                width=watermark_width,
                height=watermark_height,
                preserveAspectRatio=True,
                mask="auto",
            )

            pdf.restoreState()

        except Exception:
            pass

    # ========================================================
    # DATOS PERSONALES
    # ========================================================

    x = left_margin
    y = page_height - 50 * mm

    line_step = 8 * mm

    draw_field(
        "GRADO",
        (
            personnel.get(
                "rank_abbreviation"
            )
            or personnel.get(
                "rank"
            )
        ),
        x,
        y,
    )

    y -= line_step

    draw_field(
        "NOMBRES Y APELLIDOS",
        personnel.get(
            "full_name"
        ),
        x,
        y,
    )

    y -= line_step

    draw_field(
        "AÑO DE EGRESO",
        personnel.get(
            "graduation_year"
        ),
        x,
        y,
    )

    # ========================================================
    # PISTOLA
    # ========================================================

    y -= 11 * mm

    draw_field(
        "Nº DE PISTOLA",
        pistol.get(
            "identification_number"
        ),
        x,
        y,
    )

    y -= line_step

    draw_field(
        "MARCA",
        pistol.get(
            "manufacturer"
        ),
        x,
        y,
    )

    y -= line_step

    draw_field(
        "INDUSTRIA",
        pistol.get(
            "country"
        ),
        x,
        y,
    )

    y -= line_step

    draw_field(
        "MODELO",
        pistol.get(
            "model"
        ),
        x,
        y,
    )

    y -= line_step

    magazine_count = (
        pistol.get(
            "magazine_count",
            0,
        )
    )

    if magazine_count not in (
        None,
        "",
    ):
        try:
            magazine_count = (
                f"{int(magazine_count):02d}"
            )
        except (
            TypeError,
            ValueError,
        ):
            pass

    draw_field(
        "Nº DE CARGADORES",
        magazine_count,
        x,
        y,
    )

    y -= line_step

    draw_field(
        "OBSERVACIONES",
        (
            pistol.get(
                "observations"
            )
            or "S/O."
        ),
        x,
        y,
    )

    # ========================================================
    # CUCHILLO BAYONETA
    # ========================================================

    y -= 13 * mm

    pdf.setFont(
        "Helvetica-Bold",
        10,
    )

    pdf.drawString(
        x,
        y,
        "CUCHILLO BAYONETA",
    )

    y -= 9 * mm

    draw_field(
        "NÚMERO",
        bayonet.get(
            "identification_number"
        ),
        x,
        y,
    )

    y -= line_step

    draw_field(
        "MODELO",
        bayonet.get(
            "model"
        ),
        x,
        y,
    )

    y -= line_step

    draw_field(
        "INDUSTRIA",
        bayonet.get(
            "country"
        ),
        x,
        y,
    )

    y -= line_step

    draw_field(
        "OBSERVACIONES",
        (
            bayonet.get(
                "observations"
            )
            or "S/O."
        ),
        x,
        y,
    )

    # ========================================================
    # FOTOGRAFÍAS
    # ========================================================

    photo_top = (
        page_height - 45 * mm
    )

    top_photo_height = 47 * mm
    middle_photo_height = 34 * mm
    bottom_photo_height = 37 * mm

    # FOTO 1
    photo1_y = (
        photo_top
        - top_photo_height
    )

    draw_photo_box(
        photos[0],
        photo_x,
        photo1_y,
        photo_width,
        top_photo_height,
    )

    # FOTO 2
    photo2_y = (
        photo1_y
        - middle_photo_height
    )

    draw_photo_box(
        photos[1],
        photo_x,
        photo2_y,
        photo_width,
        middle_photo_height,
    )

    # FOTO 3
    photo3_y = (
        photo2_y
        - bottom_photo_height
    )

    draw_photo_box(
        photos[2],
        photo_x,
        photo3_y,
        photo_width,
        bottom_photo_height,
    )

    # ========================================================
    # TABLA DE FIRMAS
    # ========================================================

    table_x = left_margin
    table_width = (
        page_width
        - left_margin
        - right_margin
    )

    col_width = (
        table_width / 3
    )

    table_bottom = 34 * mm
    table_top = 102 * mm

    middle_y = 63 * mm

    # Marco exterior

    pdf.setLineWidth(0.8)

    pdf.rect(
        table_x,
        table_bottom,
        table_width,
        table_top - table_bottom,
    )

    # Columnas

    pdf.line(
        table_x + col_width,
        table_bottom,
        table_x + col_width,
        table_top,
    )

    pdf.line(
        table_x + 2 * col_width,
        table_bottom,
        table_x + 2 * col_width,
        table_top,
    )

    # División horizontal

    pdf.line(
        table_x,
        middle_y,
        table_x + table_width,
        middle_y,
    )

    # ========================================================
    # FIRMA / ACLARACIÓN / FECHA
    # ========================================================

    headers = [
        "FIRMA",
        "ACLARACIÓN",
        "FECHA",
    ]

    pdf.setFont(
        "Helvetica-Bold",
        9,
    )

    for index, header in enumerate(
        headers
    ):
        center_x = (
            table_x
            + (
                index
                * col_width
            )
            + col_width / 2
        )

        pdf.drawCentredString(
            center_x,
            middle_y + 3 * mm,
            header,
        )

    # ========================================================
    # RESPONSABLES / FIRMANTES
    # ========================================================

    signatories = report.get(
        "signatories",
        {},
    )

    verification = signatories.get(
        "verification_responsible"
    )

    logistics = signatories.get(
        "logistics_chief"
    )

    commander = signatories.get(
        "unit_commander"
    )

    signature_data = [
        {
            "person": verification,
            "role": [
                "RESPONSABLE DE LA VERIFICACIÓN",
            ],
        },
        {
            "person": logistics,
            "role": [
                "JEFE DE LA SECCIÓN IV",
                '"LOGÍSTICA"',
            ],
        },
        {
            "person": commander,
            "role": [
                "CMDTE. DEL BCA.",
            ],
        },
    ]

    for index, item in enumerate(signature_data):

        center_x = (
            table_x
            + (index * col_width)
            + col_width / 2
        )

        person = item["person"]

        # ----------------------------------------------------
        # GRADO Y NOMBRE DEL FIRMANTE
        # ----------------------------------------------------

        if person:
            rank = safe(
                person.get("rank")
            ).upper()

            full_name = safe(
                person.get("full_name")
            ).upper()

            name_y = table_bottom + 18 * mm

            pdf.setFont(
                "Helvetica-Bold",
                7,
            )

            if rank:
                pdf.drawCentredString(
                    center_x,
                    name_y + 4 * mm,
                    rank,
                )

            pdf.setFont(
                "Helvetica",
                6.5,
            )

            pdf.drawCentredString(
                center_x,
                name_y,
                full_name,
            )

        # ----------------------------------------------------
        # FUNCIÓN QUE CUMPLE EN EL DOCUMENTO
        # ----------------------------------------------------

        role_y = table_bottom + 8 * mm

        pdf.setFont(
            "Helvetica-Bold",
            6.5,
        )

        for role_line in item["role"]:
            pdf.drawCentredString(
                center_x,
                role_y,
                role_line,
            )

            role_y -= 3.5 * mm


    # ========================================================
    # DECLARACIÓN INFERIOR
    # ========================================================

    declaration_height = 9 * mm

    pdf.rect(
        table_x,
        table_bottom
        - declaration_height,
        table_width,
        declaration_height,
    )

    pdf.setFont(
        "Helvetica",
        8,
    )

    pdf.drawCentredString(
        table_x
        + table_width / 2,
        table_bottom
        - 6 * mm,
        (
            "Firmo el presente documento, como fe de la "
            "veracidad de los datos declarados."
        ),
    )

    # ========================================================
    # AUDITORÍA
    # ========================================================

    if generated_by:
        pdf.setFont(
            "Helvetica",
            5,
        )

        pdf.drawRightString(
            page_width
            - right_margin,
            8 * mm,
            (
                "Generado por FORTALEZA - "
                f"{generated_by}"
            ),
        )

    pdf.showPage()
    pdf.save()

    stream.seek(0)

    return stream